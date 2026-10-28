"""Acceptance criteria of session 01 (METODO.md, 2026-10-22), on synthetic DBFs with known truth."""
from __future__ import annotations

import datetime as dt
import re
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sinan_anotador.cli import annotate_file
from sinan_anotador.cnv import parse_def
from synthetic import make_synthetic
from dbf_writer import write_dbf

AGRAVOS = {  # Sinan NET export name -> .def (TB: version 5.0 of the form; see METODO 2026-10-22)
    "TUBENET": "TuberculNET5_0", "SIFICNET": "SifilisNET", "SIFGENET": "GestSifNET", "HANSNET": "HansNET",
    "LEIVNET": "LeishvisNET", "LTANET": "LeishtegNET", "DENGNET": "DengueNET",
}


RE_HEADER_COL = re.compile(r"\(([A-Za-z0-9_]+)\)$")


def raw_col(header: str) -> str:
    """'Sexo (CS_SEXO)' -> 'CS_SEXO'; a header without the suffix is already the raw name."""
    m = RE_HEADER_COL.search(header)
    return m.group(1) if m else header


def read_sheet(path: Path, name: str, raw: bool = True) -> tuple[list[str], list[tuple]]:
    wb = load_workbook(path, read_only=True)
    rows = list(wb[name].iter_rows(values_only=True))
    wb.close()
    header = [raw_col(h) if (raw and name == "dados") else h for h in rows[0]]
    return header, rows[1:]


def column(header, rows, name):
    i = header.index(name)
    return [r[i] for r in rows]


def norm(v):
    return "" if v is None else v


@pytest.fixture(scope="module", params=sorted(AGRAVOS))
def case(request, tab_dir, tmp_path_factory):
    stem = request.param
    tmp = tmp_path_factory.mktemp(stem)
    deffile = parse_def(tab_dir / f"{AGRAVOS[stem]}.def")
    dbf, truth = make_synthetic(deffile, tab_dir, tmp / f"{stem}.DBF", n=300, seed=sum(map(ord, stem)))
    ok, files = annotate_file(dbf, tab_dir, None, None, 1_048_575)
    assert ok, (tmp / f"{stem}_anotado.log.txt").read_text(encoding="utf-8")
    assert files == [tmp / f"{stem}_anotado.xlsx"]
    header, rows = read_sheet(files[0], "dados")
    return dict(stem=stem, deffile=deffile, truth=truth, header=header, rows=rows, files=files, tmp=tmp)


def test_c1_every_field_with_table_gets_rotulo(case):
    """Criterion 1: 100% of the DBF fields that have an available .cnv in the .def get `<campo>_rotulo`."""
    header = case["header"]
    expected = set(case["truth"]["labeled_fields"]) | ({"NU_IDADE_N"} if case["truth"]["plan"].age_band else set())
    missing = [f for f in expected if f"{f}_rotulo" not in header]
    assert not missing, missing
    assert len(expected) >= 10, (case["stem"], sorted(expected))
    for f in expected:   # label column sits right after the code column
        assert header.index(f"{f}_rotulo") - header.index(f) in (1, 2), f


def test_c2_counts_by_label_match_truth(case):
    """Criterion 2: count per label in the xlsx == truth, field by field, in 100% of the fields."""
    header, rows, truth = case["header"], case["rows"], case["truth"]["truth"]
    bad = {}
    for f, expected in truth.items():
        got = Counter(norm(v) for v in column(header, rows, f"{f}_rotulo"))
        if got != expected:
            bad[f] = (expected - got, got - expected)
    assert not bad, bad
    assert len(truth) >= 10


def test_c3_no_record_lost(case):
    assert len(case["rows"]) == case["truth"]["n"]
    labels = {norm(v) for f in case["truth"]["truth"] for v in column(case["header"], case["rows"], f"{f}_rotulo")}
    assert any(re.fullmatch(r"\[.+\] sem rótulo", lab) for lab in labels)
    assert "Ign/Branco" in labels or "Em Branco" in labels


def test_c4_age_and_dates(case):
    header, rows = case["header"], case["rows"]
    got = Counter((norm(a), b) for a, b in zip(column(header, rows, "NU_IDADE_N_rotulo"), column(header, rows, "NU_IDADE_N_anos")))
    exp = Counter((norm(a), b) for (a, b), k in case["truth"]["truth_age"].items() for _ in range(k))
    assert got == exp
    dates = column(header, rows, "DT_NOTIFIC")
    assert all(d is None or isinstance(d, (dt.date, dt.datetime)) for d in dates)
    assert [None if d is None else d.date() if isinstance(d, dt.datetime) else d for d in dates] == case["truth"]["dates"]
    assert "DT_NOTIFIC_rotulo" not in header


def municipality_names(path: Path) -> dict[str, str]:
    """Independent read of a Munic??.cnv: `<seq>  <code> <name>   <code>` -> {code: name}, name without the code."""
    names = {}
    for line in path.read_text(encoding="latin-1").splitlines():
        m = re.match(r"^\s*\d+\s+(\d{6})\s+(.*?)\s+\1\s*$", line)
        if m:
            names[m.group(1)] = " ".join(m.group(2).split())
    return names


def test_c5_municipio(case, tab_dir):
    """Criterion 5: CE municipality codes get the municipality name (CE table first); since 2026-10-25 the label is the
    name alone ("Itapiúna"), the code stays in the column to the left."""
    header, rows = case["header"], case["rows"]
    ce_names = municipality_names(tab_dir / "Municce.cnv")
    assert ce_names["230650"] == "Itapiúna" and ce_names["230440"] == "Fortaleza"
    for f in ("ID_MN_RESI", "ID_MUNICIP"):
        if f in header:
            codes = column(header, rows, f)
            labs = column(header, rows, f"{f}_rotulo")
            ce = [(str(c), l) for c, l in zip(codes, labs) if c and str(c).startswith("23") and str(c) in ce_names]
            assert ce, f
            assert all(l == ce_names[c] for c, l in ce), (f, [(c, l) for c, l in ce if l != ce_names[c]][:5])
            assert not any(l.startswith(str(c)) for c, l in zip(codes, labs) if c and l and l[0].isdigit() and len(str(c)) == 6), f


def test_long_names_in_header(case):
    """Header is 'Nome por extenso (CAMPO)'; every column keeps its raw name in parentheses; `campos` sheet lists all."""
    header, rows = read_sheet(case["files"][0], "dados", raw=False)
    assert all(RE_HEADER_COL.search(h) or h == raw_col(h) for h in header)
    named = [h for h in header if RE_HEADER_COL.search(h)]
    assert len(named) >= 0.8 * len(header), (len(named), len(header))
    assert "Sexo (CS_SEXO)" in header and "Sexo — rótulo (CS_SEXO_rotulo)" in header
    assert "Idade em anos (NU_IDADE_N_anos)" in header and "Faixa etária (NU_IDADE_N_rotulo)" in header
    assert "nome do paciente (nm_pacient)" in [h.lower() for h in header]      # form table or CSV dictionary spelling
    campos_h, campos = read_sheet(case["files"][0], "campos")
    assert campos_h == ["coluna", "nome_por_extenso", "campo_dbf", "tipo_dbf", "fonte_do_nome", "rotulo_tabwin", "tabela_de_rotulos"]
    assert [r[0] for r in campos] == [raw_col(h) for h in header]


def test_short_header_option(tab_dir, tmp_path):
    deffile = parse_def(tab_dir / "HansNET.def")
    dbf, truth = make_synthetic(deffile, tab_dir, tmp_path / "HANSNET.DBF", n=10, seed=5)
    ok, files = annotate_file(dbf, tab_dir, None, None, 1_048_575, long_names=False)
    header, rows = read_sheet(files[0], "dados", raw=False)
    assert "CS_SEXO" in header and not any("(" in h for h in header)


def test_csv_dictionary_overrides_names(tab_dir, tmp_path):
    (tmp_path / "dic").mkdir()
    (tmp_path / "dic" / "dicionario_sifg.csv").write_text("variavel,variavel_dbf_observado,nome_campo\nCS_SEXO,CS_SEXO,Sexo do paciente\n"
                                              "TPEVIDENCI,TPEVIDENCI,Classificação clínica da sífilis\n", encoding="utf-8")
    deffile = parse_def(tab_dir / "GestSifNET.def")
    dbf, truth = make_synthetic(deffile, tab_dir, tmp_path / "SIFGENET.DBF", n=10, seed=5)
    ok, files = annotate_file(dbf, tab_dir, None, None, 1_048_575, csv_dir=tmp_path / "dic")
    header, rows = read_sheet(files[0], "dados", raw=False)
    assert "Sexo do paciente (CS_SEXO)" in header and "Classificação clínica da sífilis (TPEVIDENCI)" in header
    campos_h, campos = read_sheet(files[0], "campos")
    assert any(r[0] == "CS_SEXO" and r[4] == "dicionário CSV (dicionario_sifg.csv)" for r in campos)


def test_dictionary_sheet(case):
    header, rows = read_sheet(case["files"][0], "dicionario")
    assert header == ["campo", "tabela", "codigos", "rotulo", "uso"]
    fields = {r[0] for r in rows}
    assert set(case["truth"]["labeled_fields"]) <= fields
    assert any(r[1] == "Municbr.cnv" and r[4] == "rótulo (complemento)" for r in rows)
    assert any(r[1] == "Municce.cnv" for r in rows)


def test_log_file(case):
    log = (case["tmp"] / f"{case['stem']}_anotado.log.txt").read_text(encoding="utf-8")
    assert "RESULTADO: OK" in log and AGRAVOS[case["stem"]] in log and "cp850" in log


def test_c6_dependencies_only_packaged():
    """Criterion 6: only pandas, numpy, openpyxl and dbfread besides the standard library."""
    import ast
    import sys
    allowed = {"pandas", "numpy", "openpyxl", "dbfread"}
    stdlib = set(sys.stdlib_module_names)
    src = Path(__file__).resolve().parents[1] / "src" / "sinan_anotador"
    imported = set()
    for p in src.glob("*.py"):
        for node in ast.walk(ast.parse(p.read_text(encoding="utf-8"))):
            if isinstance(node, ast.Import):
                imported.update(a.name.split(".")[0] for a in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module and node.level == 0:
                imported.add(node.module.split(".")[0])
    third = {m for m in imported if m not in stdlib and m not in {"sinan_anotador"}}
    assert third <= allowed, third


def test_c7_split_by_year(tab_dir, tmp_path):
    """Criterion 7: above the row limit the output is split by notification year, with a warning."""
    deffile = parse_def(tab_dir / "GestSifNET.def")
    dbf, truth = make_synthetic(deffile, tab_dir, tmp_path / "SIFGENET.DBF", n=25, seed=7, years=(2019, 2020, 2021))
    ok, files = annotate_file(dbf, tab_dir, None, None, max_data_rows=8)
    assert ok
    names = sorted(f.name for f in files)
    assert all(re.fullmatch(r"SIFGENET_anotado_(\d{4}|sem_ano)(_parte\d+)?\.xlsx", n) for n in names), names
    total = 0
    years = Counter(d.year if d else None for d in truth["dates"])
    for f in files:
        header, rows = read_sheet(f, "dados")
        assert len(rows) <= 8
        total += len(rows)
        m = re.search(r"_(\d{4})", f.name)
        if m:
            assert all(d.year == int(m.group(1)) for d in column(header, rows, "DT_NOTIFIC") if d)
    assert total == 25
    log = (tmp_path / "SIFGENET_anotado.log.txt").read_text(encoding="utf-8")
    assert "AVISO" in log and "dividida por ano" in log


def test_encoding_detection_cp1252(tab_dir, tmp_path):
    deffile = parse_def(tab_dir / "HansNET.def")
    dbf, truth = make_synthetic(deffile, tab_dir, tmp_path / "HANSNET.DBF", n=40, seed=3, encoding="cp1252")
    ok, files = annotate_file(dbf, tab_dir, None, None, 1_048_575)
    header, rows = read_sheet(files[0], "dados")
    assert column(header, rows, "NM_PACIENT") == [r["NM_PACIENT"] for r in truth["records"]]
    assert "cp1252" in (tmp_path / "HANSNET_anotado.log.txt").read_text(encoding="utf-8").split("Codificação")[1].split("\n")[0]


def test_deleted_records_are_not_exported(tab_dir, tmp_path):
    fields = [("NU_NOTIFIC", "C", 7, 0), ("DT_NOTIFIC", "D", 8, 0), ("CS_SEXO", "C", 1, 0)]
    recs = [dict(NU_NOTIFIC=f"{i:07d}", DT_NOTIFIC=dt.date(2020, 1, 1), CS_SEXO="M") for i in range(5)]
    dbf = write_dbf(tmp_path / "HANSNET.DBF", fields, recs, deleted={2})
    ok, files = annotate_file(dbf, tab_dir, None, None, 1_048_575)
    header, rows = read_sheet(files[0], "dados")
    assert len(rows) == 4
    assert "apagados" in (tmp_path / "HANSNET_anotado.log.txt").read_text(encoding="utf-8")
