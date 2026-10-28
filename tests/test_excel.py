"""Writing the xlsx: control characters that openpyxl refuses are removed and reported."""
import datetime as dt

import pandas as pd
from openpyxl import load_workbook

from sinan_anotador.excel import _cell_value, cleaned_note, write_xlsx


def test_cell_value_strips_control_characters():
    assert _cell_value("PACIENTE FEZ TRATAMENTO\x00\x1a") == ("PACIENTE FEZ TRATAMENTO", True)
    assert _cell_value("linha 1\nlinha 2\tcol\r") == ("linha 1\nlinha 2\tcol\r", False)   # tab, LF and CR are legal
    assert _cell_value("sem nada") == ("sem nada", False)
    assert _cell_value(None) == (None, False)
    assert _cell_value(float("nan")) == (None, False)
    assert _cell_value(pd.NA) == (None, False)
    assert _cell_value(pd.Timestamp("2024-07-12")) == (dt.datetime(2024, 7, 12), False)
    assert _cell_value(12) == (12, False)


def test_write_xlsx_survives_control_characters_and_logs_them(tmp_path):
    data = pd.DataFrame({
        "NU_NOTIFIC": ["0000001", "0000002", "0000003"],
        "DS_OBS": ["RN C VDRL 1/2 INICIOU TRATAMENTO .\x00", "sem controle", "\x1aTTO\x07"],
        "NM_PACIENT": ["ANA\x00", "BIA", "CLARA"],
        "DT_NOTIFIC": [dt.date(2024, 9, 11), dt.date(2024, 9, 12), None],
    })
    dictionary = pd.DataFrame({"campo": ["DS_OBS"], "rotulo": ["obs\x00"]})
    files, notes = write_xlsx(data, dictionary, tmp_path / "SIFICNET_anotado")
    assert [f.name for f in files] == ["SIFICNET_anotado.xlsx"]
    assert len(notes) == 1 and notes[0].startswith("AVISO: 4 células continham caracteres de controle"), notes
    assert "dados.DS_OBS (2)" in notes[0] and "dados.NM_PACIENT (1)" in notes[0] and "dicionario.rotulo (1)" in notes[0]
    ws = load_workbook(files[0], read_only=True)["dados"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("NU_NOTIFIC", "DS_OBS", "NM_PACIENT", "DT_NOTIFIC")
    assert rows[1][1] == "RN C VDRL 1/2 INICIOU TRATAMENTO ." and rows[1][2] == "ANA"
    assert rows[3][1] == "TTO"
    assert isinstance(rows[1][3], dt.datetime) and rows[1][3].date() == dt.date(2024, 9, 11)


def test_write_xlsx_no_note_when_clean(tmp_path):
    data = pd.DataFrame({"A": ["x", "y"], "B": [1, 2]})
    files, notes = write_xlsx(data, pd.DataFrame({"campo": ["A"]}), tmp_path / "limpo")
    assert notes == [] and files[0].exists()
    assert cleaned_note({}) is None
